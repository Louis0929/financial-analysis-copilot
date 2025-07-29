"""
Financial Statement Analysis Co-Pilot
Main application script for analyzing financial reports using Google Gemini Pro API
Supports multiple file formats: TXT, PDF, DOCX, XLSX, CSV

Author: [Your Name]
Date: [Current Date]
Purpose: University project for automated financial statement analysis
"""

import os
import sys
import pandas as pd
from pathlib import Path
from config import gemini_model
from prompts import FINANCIAL_ANALYSIS_PROMPT

# Import libraries for different file formats
try:
    import PyPDF2
    import docx
    from openpyxl import load_workbook
except ImportError as e:
    print(f"⚠️ Some file format libraries are missing: {e}")
    print("Install with: pip install -r requirements.txt")

def read_txt_file(filepath):
    """Read content from a TXT file"""
    with open(filepath, 'r', encoding='utf-8') as file:
        return file.read()

def read_pdf_file(filepath):
    """Read content from a PDF file"""
    try:
        import PyPDF2
        content = ""
        with open(filepath, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                content += page.extract_text() + "\n"
        return content
    except ImportError:
        raise ImportError("PyPDF2 is required for PDF files. Install with: pip install PyPDF2")

def read_docx_file(filepath):
    """Read content from a Word document"""
    try:
        import docx
        doc = docx.Document(filepath)
        content = ""
        for paragraph in doc.paragraphs:
            content += paragraph.text + "\n"
        return content
    except ImportError:
        raise ImportError("python-docx is required for Word files. Install with: pip install python-docx")

def read_excel_file(filepath):
    """Read content from an Excel file"""
    try:
        # Read all sheets and combine content
        excel_file = pd.ExcelFile(filepath)
        content = ""
        
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(filepath, sheet_name=sheet_name)
            content += f"\n--- SHEET: {sheet_name} ---\n"
            content += df.to_string(index=False) + "\n"
        
        return content
    except ImportError:
        raise ImportError("pandas and openpyxl are required for Excel files. Install with: pip install pandas openpyxl")

def read_csv_file(filepath):
    """Read content from a CSV file"""
    try:
        df = pd.read_csv(filepath)
        return df.to_string(index=False)
    except ImportError:
        raise ImportError("pandas is required for CSV files. Install with: pip install pandas")

def read_report(filepath):
    """
    Read financial report content from various file formats
    Supports: TXT, PDF, DOCX, XLSX, CSV
    
    Args:
        filepath (str): Path to the financial report file
        
    Returns:
        str: Content of the financial report
    """
    try:
        # Check if file exists
        if not os.path.exists(filepath):
            print(f"❌ Error: File not found at {filepath}")
            return None
        
        # Get file extension
        file_extension = Path(filepath).suffix.lower()
        
        # Read based on file type
        if file_extension == '.txt':
            content = read_txt_file(filepath)
        elif file_extension == '.pdf':
            content = read_pdf_file(filepath)
        elif file_extension == '.docx':
            content = read_docx_file(filepath)
        elif file_extension in ['.xlsx', '.xls']:
            content = read_excel_file(filepath)
        elif file_extension == '.csv':
            content = read_csv_file(filepath)
        else:
            print(f"❌ Unsupported file format: {file_extension}")
            print("📁 Supported formats: .txt, .pdf, .docx, .xlsx, .xls, .csv")
            return None
        
        print(f"✅ Successfully loaded {file_extension.upper()} report from: {filepath}")
        print(f"📄 Content length: {len(content)} characters")
        return content
        
    except Exception as e:
        print(f"❌ Error reading file: {e}")
        return None

def find_financial_file():
    """
    Look for financial report files in the data directory
    Returns the first found file or None
    """
    data_dir = "data"
    if not os.path.exists(data_dir):
        return None
    
    # Supported file extensions
    supported_extensions = ['.txt', '.pdf', '.docx', '.xlsx', '.xls', '.csv']
    
    for file in os.listdir(data_dir):
        file_path = os.path.join(data_dir, file)
        if os.path.isfile(file_path) and Path(file).suffix.lower() in supported_extensions:
            return file_path
    
    return None

def analyze_financial_report(report_text):
    """
    Send the financial report to Gemini for analysis
    
    Args:
        report_text (str): The financial report content to analyze
        
    Returns:
        str: AI-generated analysis or None if error occurred
    """
    try:
        # Check if model is available
        if gemini_model is None:
            print("❌ Gemini model not initialized. Check your API key configuration.")
            return None
        
        # Format the prompt with the report text
        formatted_prompt = FINANCIAL_ANALYSIS_PROMPT.format(report_text=report_text)
        
        print("🔄 Sending report to Gemini for analysis...")
        
        # Generate analysis using Gemini
        response = gemini_model.generate_content(formatted_prompt)
        
        print("✅ Analysis completed successfully!")
        return response.text
        
    except Exception as e:
        print(f"❌ Error during analysis: {e}")
        if "quota" in str(e).lower():
            print("💡 Tip: Try again in a few minutes or use gemini-1.5-flash for better rate limits")
        return None

def display_analysis(analysis):
    """
    Display the analysis results in a formatted way
    
    Args:
        analysis (str): The AI-generated analysis to display
    """
    print("\n" + "="*80)
    print("FINANCIAL STATEMENT ANALYSIS RESULTS")
    print("="*80)
    print(analysis)
    print("="*80)

def main():
    """
    Main function to orchestrate the financial analysis process
    """
    print("🚀 Financial Statement Analysis Co-Pilot Starting...")
    print("-" * 60)
    
    # Step 1: Look for financial report file
    print("📁 Looking for financial report files...")
    
    # First try the default sample file
    report_filepath = os.path.join("data", "sample_report.txt")
    
    # If not found, look for any supported file
    if not os.path.exists(report_filepath):
        found_file = find_financial_file()
        if found_file:
            report_filepath = found_file
            print(f"📄 Found financial report: {report_filepath}")
        else:
            print("❌ No financial report files found in data/ directory")
            print("📁 Please add a file in one of these formats: .txt, .pdf, .docx, .xlsx, .csv")
            sys.exit(1)
    
    # Step 2: Read the financial report content
    print("📖 Reading financial report...")
    report_content = read_report(report_filepath)
    
    if report_content is None:
        print("❌ Cannot proceed without report content. Exiting.")
        sys.exit(1)
    
    # Step 3: Analyze the report using Gemini AI
    print("🤖 Analyzing financial report with AI...")
    analysis_result = analyze_financial_report(report_content)
    
    if analysis_result is None:
        print("❌ Analysis failed. Please check your configuration and try again.")
        sys.exit(1)
    
    # Step 4: Display the results
    display_analysis(analysis_result)
    
    print("\n✅ Financial analysis completed successfully!")
    print("💡 Tips:")
    print("   • Place your financial reports in the 'data/' folder")
    print("   • Supported formats: TXT, PDF, DOCX, XLSX, CSV")
    print("   • The app will automatically detect and process the file")

if __name__ == "__main__":
    main() 